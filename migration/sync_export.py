"""Build an editable Excel copy of the live tracker.

Usage:
    python3 -I migration/sync_export.py <live_dir> <out.xlsx>

<live_dir> is a database export: one folder per collection with one JSON
file per record, plus versions.json ({"tasks/kndz-001": 2, ...}).

The workbook carries a hidden _sync sheet with what each row looked like at
export time. sync_import.py uses it to apply only the cells someone changed,
and to spot rows that were also changed in the app since the export.
"""
import datetime
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

LIVE, OUT = sys.argv[1], sys.argv[2]


def load(coll):
    p = os.path.join(LIVE, coll)
    if not os.path.isdir(p):
        return {}
    return {f[:-5]: json.load(open(os.path.join(p, f))) for f in sorted(os.listdir(p)) if f.endswith(".json")}


VERS = json.load(open(os.path.join(LIVE, "versions.json")))
brands, tasks, sow, access = load("brands"), load("tasks"), load("sow"), load("access")
bps, bsteps, steps = load("bps"), load("brand_steps"), load("steps")
cfg = load("config")
lists, people = cfg["lists"], [p["name"] for p in cfg["people"]["people"] if p.get("active", True)]
DEPTS = lists.get("department") or ["Other"]
border = sorted(brands, key=lambda b: (brands[b].get("order", 999), brands[b]["name"]))
bname = {b: brands[b]["name"] for b in brands}

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1D5C8A")
RO_FILL = PatternFill("solid", fgColor="ECE9E2")     # grey = worked out for you, edits ignored
ID_FONT = Font(color="8B9098", size=9)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

wb = Workbook()
sync_rows = []          # (sheet, id, version, json of exported values)


def iso_to_date(s):
    try:
        return datetime.datetime.strptime(s, "%Y-%m-%d").date() if s else None
    except ValueError:
        return None


def sheet(title, cols, widths, ro=()):
    ws = wb.create_sheet(title)
    ws.append(cols)
    for i, c in enumerate(cols, 1):
        cell = ws.cell(1, i)
        cell.font, cell.fill = HEAD, HEAD_FILL if c not in ro else PatternFill("solid", fgColor="5F6670")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = "C2"
    ws.row_dimensions[1].height = 32
    return ws


def finish(ws, ncols, ro_cols, date_cols=(), wrap_cols=()):
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.alignment = WRAP if c.column in wrap_cols else TOP
            if c.column in ro_cols:
                c.fill = RO_FILL
            if c.column in date_cols:
                c.number_format = "ddd d mmm yyyy"
        row[0].font = ID_FONT
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{max(ws.max_row, 2)}"


# ---------------------------------------------------------------- Lists (hidden, feeds the dropdowns)
wl = wb.create_sheet("Lists")
named = {
    "Brands": [bname[b] for b in border], "Teams": DEPTS, "People": sorted(people),
    "TaskStatus": lists["task_status"], "Priority": lists["priority"], "Stage": lists["stage"],
    "SowStage": lists["sow_stage"], "AccessStatus": lists["access_status"], "AccessLevel": lists["access_level"],
    "Platform": lists["platform"], "StepStatus": lists["checklist"], "BpStatus": lists["bp_status"], "YesNo": ["Yes", "No"],
}
ref = {}
for i, (name, vals) in enumerate(named.items(), 1):
    col = get_column_letter(i)
    wl.cell(1, i, name)
    for j, v in enumerate(vals, 2):
        wl.cell(j, i, v)
    ref[name] = f"Lists!${col}$2:${col}${max(len(vals) + 1, 2)}"
wl.sheet_state = "hidden"


def dropdown(ws, col, name, strict=True, last=1500):
    dv = DataValidation(type="list", formula1=ref[name], allow_blank=True, showErrorMessage=strict)
    if strict:
        dv.error, dv.errorTitle = "Pick a value from the list.", "Not in the list"
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{last}")


def due_formula(r, date_col, done_test):
    return (f'=IF(OR({done_test},NOT(ISNUMBER({date_col}{r}))),"",IF({date_col}{r}<TODAY(),"Overdue",'
            f'IF({date_col}{r}-TODAY()<=3,"Due Soon","On Time")))')


# ---------------------------------------------------------------- Read me
rm = wb.active
rm.title = "Read me"
guide = [
    ("Valerie COE Tracker: Excel copy", ""),
    ("Exported", datetime.datetime.now().strftime("%d %b %Y %H:%M")),
    ("", ""),
    ("How to use", ""),
    ("1", "Change any white cell. Use the dropdowns for status, owner, team and dates."),
    ("2", "Grey cells and the ID column are filled in for you. Changes there are ignored."),
    ("3", "To add a task, fill a new row at the bottom of Tasks and leave ID empty. Brand and Task are required."),
    ("4", "To finish a task, set Status to Complete. Do not delete rows. Deleted rows are not removed from the tracker."),
    ("5", "Blocked needs a reason in 'Blocked because' and a person in 'Waiting on'."),
    ("6", "Onboarding Steps: one column per brand. Pick Done, In Progress, Not Started or N/A."),
    ("7", "Send the file back. Only the cells you changed are written to the tracker."),
    ("8", "If someone changed the same cell in the tracker after this export, the tracker keeps their value and you get a list of those cells."),
    ("", ""),
    ("Sheets", "Tasks, Brands, SOW, Platform Access, Onboarding Steps, Best Practices"),
]
for a, b in guide:
    rm.append([a, b])
rm["A1"].font = Font(bold=True, size=14)
rm["A4"].font = rm["A14"].font = Font(bold=True)
rm.column_dimensions["A"].width = 14
rm.column_dimensions["B"].width = 110

# ---------------------------------------------------------------- Tasks
TCOLS = ["ID", "Brand", "Team", "Task", "Owner", "Support team", "Status", "Priority", "Start", "Deadline",
         "Due flag", "Notes / next step", "Blocked because", "Waiting on", "Brief link", "Creative link"]
ws = sheet("Tasks", TCOLS, [14, 16, 16, 44, 13, 22, 13, 10, 13, 13, 11, 60, 30, 13, 18, 18], ro=("ID", "Due flag"))
dorder = {d: i for i, d in enumerate(DEPTS)}


def urgency(t):
    if t.get("status") == "Complete":
        return 9
    if t.get("status") == "Blocked":
        return 0
    return 1 if t.get("deadline") else 2


rows = sorted(tasks.items(), key=lambda kv: (border.index(kv[1]["brand"]) if kv[1].get("brand") in border else 999,
                                             dorder.get(kv[1].get("department"), 99), urgency(kv[1]), kv[1].get("deadline") or "9999"))
for r, (tid, t) in enumerate(rows, 2):
    blk = t.get("blocker") or {}
    vals = {"Brand": bname.get(t.get("brand"), ""), "Team": t.get("department") or "Other", "Task": t.get("title", ""),
            "Owner": t.get("owner", ""), "Support team": ", ".join(t.get("support") or []), "Status": t.get("status", ""),
            "Priority": t.get("priority", ""), "Start": t.get("start", ""), "Deadline": t.get("deadline", ""),
            "Notes / next step": t.get("notes", ""), "Blocked because": blk.get("text", ""), "Waiting on": blk.get("person", ""),
            "Brief link": t.get("brief_link", ""), "Creative link": t.get("creative_link", "")}
    ws.append([tid, vals["Brand"], vals["Team"], vals["Task"], vals["Owner"], vals["Support team"], vals["Status"], vals["Priority"],
               iso_to_date(vals["Start"]), iso_to_date(vals["Deadline"]), due_formula(r, "J", f'D{r}="",G{r}="Complete"'),
               vals["Notes / next step"], vals["Blocked because"], vals["Waiting on"], vals["Brief link"], vals["Creative link"]])
    sync_rows.append(("Tasks", tid, VERS.get("tasks/" + tid, 0), vals))
for r in range(ws.max_row + 1, ws.max_row + 201):          # empty rows ready for new tasks
    ws.cell(r, 11, due_formula(r, "J", f'D{r}="",G{r}="Complete"'))
finish(ws, len(TCOLS), ro_cols=(1, 11), date_cols=(9, 10), wrap_cols=(4, 12, 13))
for col, name in [("B", "Brands"), ("C", "Teams"), ("G", "TaskStatus"), ("H", "Priority")]:
    dropdown(ws, col, name)
for col in ("E", "N"):
    dropdown(ws, col, "People", strict=False)

# ---------------------------------------------------------------- Brands
BCOLS = ["ID", "Brand", "Founder", "Website", "Stage", "Brand lead", "Sprint", "Paused", "Next step", "Open tasks", "Blocked", "Late"]
ws = sheet("Brands", BCOLS, [16, 20, 16, 30, 20, 13, 30, 9, 70, 10, 9, 8], ro=("ID", "Open tasks", "Blocked", "Late"))
for r, b in enumerate(border, 2):
    d = brands[b]
    vals = {"Brand": d["name"], "Founder": d.get("founder", ""), "Website": d.get("website", ""), "Stage": d.get("stage", ""),
            "Brand lead": d.get("lead", ""), "Sprint": d.get("sprint_label", ""), "Paused": "Yes" if d.get("paused") else "No",
            "Next step": d.get("next_step", "")}
    ws.append([b] + list(vals.values()) + [
        f'=COUNTIFS(Tasks!$B:$B,$B{r},Tasks!$D:$D,"<>",Tasks!$G:$G,"<>Complete")',
        f'=COUNTIFS(Tasks!$B:$B,$B{r},Tasks!$G:$G,"Blocked")',
        f'=COUNTIFS(Tasks!$B:$B,$B{r},Tasks!$K:$K,"Overdue")'])
    sync_rows.append(("Brands", b, VERS.get("brands/" + b, 0), vals))
finish(ws, len(BCOLS), ro_cols=(1, 10, 11, 12), wrap_cols=(9,))
dropdown(ws, "E", "Stage")
dropdown(ws, "H", "YesNo")
dropdown(ws, "F", "People", strict=False)

# ---------------------------------------------------------------- SOW
SCOLS = ["ID", "Brand", "SOW stage", "Owner", "Support team", "Start date", "Draft due", "Target sign", "Due flag",
         "Blocker / dependency", "Next step", "SOW link", "Growth Forecast link", "GF updated", "Batch"]
ws = sheet("SOW", SCOLS, [16, 18, 18, 11, 20, 13, 13, 13, 11, 40, 50, 22, 22, 10, 10], ro=("ID", "Due flag"))
srows = sorted(sow.items(), key=lambda kv: (kv[1].get("target_sign") or "9999", bname.get(kv[0], "")))
for r, (sid, s) in enumerate(srows, 2):
    vals = {"Brand": bname.get(sid, sid), "SOW stage": s.get("stage", ""), "Owner": s.get("owner", ""),
            "Support team": ", ".join(s.get("support") or []), "Start date": s.get("start_date", ""), "Draft due": s.get("draft_due", ""),
            "Target sign": s.get("target_sign", ""), "Blocker / dependency": s.get("blocker", ""), "Next step": s.get("next_step", ""),
            "SOW link": s.get("sow_link", ""), "Growth Forecast link": s.get("gf_link", ""),
            "GF updated": "Yes" if s.get("gf_updated") else "No", "Batch": s.get("batch", "")}
    ws.append([sid, vals["Brand"], vals["SOW stage"], vals["Owner"], vals["Support team"], iso_to_date(vals["Start date"]),
               iso_to_date(vals["Draft due"]), iso_to_date(vals["Target sign"]),
               f'=IF(C{r}="Signed","",IF(ISNUMBER(H{r}),IF(H{r}<TODAY(),"Overdue",IF(H{r}-TODAY()<=3,"Due Soon","On Time")),IF(ISNUMBER(G{r}),IF(G{r}<TODAY(),"Overdue",IF(G{r}-TODAY()<=3,"Due Soon","On Time")),"")))',
               vals["Blocker / dependency"], vals["Next step"], vals["SOW link"], vals["Growth Forecast link"], vals["GF updated"], vals["Batch"]])
    sync_rows.append(("SOW", sid, VERS.get("sow/" + sid, 0), vals))
finish(ws, len(SCOLS), ro_cols=(1, 2, 9), date_cols=(6, 7, 8), wrap_cols=(10, 11))
dropdown(ws, "C", "SowStage")
dropdown(ws, "N", "YesNo")
dropdown(ws, "D", "People", strict=False)

# ---------------------------------------------------------------- Platform access
ACOLS = ["ID", "Brand", "Platform", "Access level", "Needed for", "Valerie owner", "Brand contact", "Requested on",
         "Deadline", "Status", "Days waiting", "Blocker", "Notes"]
ws = sheet("Platform Access", ACOLS, [16, 18, 22, 18, 16, 13, 14, 13, 13, 13, 10, 36, 40], ro=("ID", "Days waiting"))
for r, (aid, a) in enumerate(sorted(access.items(), key=lambda kv: bname.get(kv[1]["brand"], "")), 2):
    vals = {"Brand": bname.get(a["brand"], ""), "Platform": a.get("platform", ""), "Access level": a.get("access_level", ""),
            "Needed for": a.get("needed_for", ""), "Valerie owner": a.get("owner", ""), "Brand contact": a.get("brand_contact", ""),
            "Requested on": a.get("requested_date", ""), "Deadline": a.get("deadline", ""), "Status": a.get("status", ""),
            "Blocker": a.get("blocker", ""), "Notes": a.get("notes", "")}
    ws.append([aid, vals["Brand"], vals["Platform"], vals["Access level"], vals["Needed for"], vals["Valerie owner"], vals["Brand contact"],
               iso_to_date(vals["Requested on"]), iso_to_date(vals["Deadline"]), vals["Status"],
               f'=IF(OR(NOT(ISNUMBER(H{r})),J{r}="Granted",J{r}="Not Needed"),"",TODAY()-H{r})', vals["Blocker"], vals["Notes"]])
    sync_rows.append(("Platform Access", aid, VERS.get("access/" + aid, 0), vals))
finish(ws, len(ACOLS), ro_cols=(1, 11), date_cols=(8, 9), wrap_cols=(12, 13))
for col, name in [("B", "Brands"), ("C", "Platform"), ("D", "AccessLevel"), ("J", "AccessStatus")]:
    dropdown(ws, col, name)
dropdown(ws, "F", "People", strict=False)

# ---------------------------------------------------------------- Onboarding steps (step x brand grid)
step_list = sorted(steps.values(), key=lambda s: s["order"])
OCOLS = ["Code", "Phase", "Step", "Gate"] + [bname[b] for b in border]
ws = sheet("Onboarding Steps", OCOLS, [8, 26, 52, 7] + [13] * len(border), ro=("Code", "Phase", "Step", "Gate"))
ws.freeze_panes = "E2"
for s in step_list:
    row = [s["code"], f'{s["phase"]} {s["phase_name"]}', s["name"], "Gate" if s.get("is_gate") else ""]
    for b in border:
        row.append(((bsteps.get(b, {}).get("s") or {}).get(s["code"]) or {}).get("status", ""))
    ws.append(row)
for b in border:
    sync_rows.append(("Onboarding Steps", b, VERS.get("brand_steps/" + b, 0),
                      {s["code"]: ((bsteps.get(b, {}).get("s") or {}).get(s["code"]) or {}).get("status", "") for s in step_list}))
finish(ws, len(OCOLS), ro_cols=(1, 2, 3, 4), wrap_cols=(3,))
for i in range(5, len(OCOLS) + 1):
    dropdown(ws, get_column_letter(i), "StepStatus", last=len(step_list) + 1)

# ---------------------------------------------------------------- Best Practices
secs = lists["bp_section"]
PCOLS = ["ID", "Brand", "BP owner", "Due date"] + secs + ["Blocker", "BP document link", "Notes"]
ws = sheet("Best Practices", PCOLS, [16, 18, 12, 13] + [13] * len(secs) + [30, 22, 40], ro=("ID",))
for bid in [b for b in border if b in bps]:
    d = bps[bid]
    vals = {"Brand": bname[bid], "BP owner": d.get("owner", ""), "Due date": d.get("due_date", "")}
    vals.update({s: (d.get("sections") or {}).get(s, "Not Started") for s in secs})
    vals.update({"Blocker": d.get("blocker", ""), "BP document link": d.get("doc_link", ""), "Notes": d.get("notes", "")})
    ws.append([bid, vals["Brand"], vals["BP owner"], iso_to_date(vals["Due date"])] + [vals[s] for s in secs] +
              [vals["Blocker"], vals["BP document link"], vals["Notes"]])
    sync_rows.append(("Best Practices", bid, VERS.get("bps/" + bid, 0), vals))
finish(ws, len(PCOLS), ro_cols=(1, 2), date_cols=(4,), wrap_cols=(len(PCOLS),))
for i in range(5, 5 + len(secs)):
    dropdown(ws, get_column_letter(i), "BpStatus")
dropdown(ws, "C", "People", strict=False)

# ---------------------------------------------------------------- _sync (very hidden)
sy = wb.create_sheet("_sync")
sy.append(["sheet", "id", "version", "values"])
for sh, rid, ver, vals in sync_rows:
    sy.append([sh, rid, ver, json.dumps(vals, ensure_ascii=False)])
sy.sheet_state = "veryHidden"

wb.move_sheet("Tasks", offset=-(wb.sheetnames.index("Tasks") - 1))
wb.save(OUT)
print(f"saved {OUT}: tasks={len(tasks)} brands={len(brands)} sow={len(sow)} access={len(access)} bps={len(bps)} steps={len(step_list)}")

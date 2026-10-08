"""Read an edited Excel copy and work out the tracker updates.

Usage:
    python3 -I migration/sync_import.py <edited.xlsx> <live_dir> <out_dir>

<live_dir> is a fresh database export taken just before the import (same
shape as for sync_export.py, with versions.json). Nothing is written to the
tracker by this script. It writes:
    <out_dir>/docs/...       one JSON file per record to update or create
    <out_dir>/batch_N.json   database batches of at most 50 writes, pinned to
                             the live version so a newer app edit is never overwritten
    <out_dir>/report.md      every change, every skipped cell and why
Rule: a cell is applied only when it differs from the export snapshot. If the
same field was also changed in the tracker since the export, the tracker
value wins and the cell is listed as a conflict.
"""
import datetime
import json
import os
import re
import sys

import openpyxl

XLSX, LIVE, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
wb = openpyxl.load_workbook(XLSX)


def load(coll):
    p = os.path.join(LIVE, coll)
    return {f[:-5]: json.load(open(os.path.join(p, f))) for f in os.listdir(p) if f.endswith(".json")} if os.path.isdir(p) else {}


VERS = json.load(open(os.path.join(LIVE, "versions.json")))
live = {c: load(c) for c in ("brands", "tasks", "sow", "access", "bps", "brand_steps", "steps", "config")}
lists = live["config"]["lists"]
people = {p["name"] for p in live["config"]["people"]["people"]}
by_name = {b["name"].lower(): bid for bid, b in live["brands"].items()}

base = {}
for sh, rid, ver, vals in wb["_sync"].iter_rows(min_row=2, values_only=True):
    base[(sh, rid)] = (ver, json.loads(vals))

writes = {}       # (collection, id) -> {"op": "update"/"set", "data": {...}}
report = {"applied": [], "skipped": [], "conflicts": [], "new": [], "missing": []}


def txt(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).replace("\xa0", " ").strip()


def date_val(v, where):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    s = txt(v)
    if not s or re.match(r"^\d{4}-\d{2}-\d{2}$", s):
        return s
    report["skipped"].append((where, f"'{s}' is not a date. Use the date picker or type 31/10/2026"))
    return None


def rows(ws):
    head = [txt(c.value) for c in ws[1]]
    for r in ws.iter_rows(min_row=2):
        vals = {head[i]: c.value for i, c in enumerate(r) if i < len(head) and head[i]}
        if any(txt(v) for k, v in vals.items() if k not in ("Due flag", "Days waiting", "Open tasks", "Blocked", "Late")):
            yield r[0].row, vals


def queue(coll, rid, patch, op="update"):
    w = writes.setdefault((coll, rid), {"op": op, "data": {}})
    for k, v in patch.items():
        if isinstance(v, dict) and isinstance(w["data"].get(k), dict):
            w["data"][k].update(v)
        else:
            w["data"][k] = v


def check(value, allowed, where, field):
    if value and value not in allowed:
        report["skipped"].append((where, f"{field} '{value}' is not in the list"))
        return False
    return True


def apply(sheet, coll, rid, row_no, edited, to_db, live_doc, label):
    """edited: {column: value as text}. to_db: {column: (field, convert)}."""
    snap = base.get((sheet, rid))
    if snap is None:
        report["skipped"].append((f"{sheet} row {row_no}", f"ID {rid} was not in the export. Leave the ID cell as it was"))
        return
    _, before = snap
    live_ver = VERS.get(f"{coll}/{rid}", 0)
    for col, now in edited.items():
        if col not in to_db or now is None:
            continue
        old = before.get(col, "")
        if now == old:
            continue
        field, conv = to_db[col]
        new_val = conv(now)
        live_val = live_doc.get(field)
        live_txt = ", ".join(live_val) if isinstance(live_val, list) else ("Yes" if live_val is True else "No" if live_val is False and col in ("Paused", "GF updated") else txt(live_val))
        if live_ver != snap[0] and live_txt != old:
            report["conflicts"].append((f"{sheet} row {row_no}", f"{label}: {col}. Tracker now says '{live_txt}', your file says '{now}'. Kept the tracker value"))
            continue
        queue(coll, rid, {field: new_val})
        report["applied"].append((f"{sheet} row {row_no}", f"{label}: {col} '{old}' to '{now}'"))


ident = lambda v: v
split = lambda v: [p.strip() for p in v.split(",") if p.strip()]
yes = lambda v: v == "Yes"

# ---------------------------------------------------------------- Tasks
seen = set()
next_n = {}
for tid in live["tasks"]:
    m = re.match(r"(.+)-(\d+)$", tid)
    if m:
        next_n[m.group(1)] = max(next_n.get(m.group(1), 0), int(m.group(2)))
for r, v in rows(wb["Tasks"]):
    where = f"Tasks row {r}"
    rid = txt(v.get("ID"))
    brand_txt = txt(v.get("Brand"))
    brand = by_name.get(brand_txt.lower())
    e = {"Brand": brand_txt, "Team": txt(v.get("Team")), "Task": txt(v.get("Task")), "Owner": txt(v.get("Owner")),
         "Support team": txt(v.get("Support team")), "Status": txt(v.get("Status")), "Priority": txt(v.get("Priority")),
         "Start": date_val(v.get("Start"), where), "Deadline": date_val(v.get("Deadline"), where),
         "Notes / next step": txt(v.get("Notes / next step")), "Blocked because": txt(v.get("Blocked because")),
         "Waiting on": txt(v.get("Waiting on")), "Brief link": txt(v.get("Brief link")), "Creative link": txt(v.get("Creative link"))}
    if brand_txt and not brand:
        report["skipped"].append((where, f"Brand '{brand_txt}' not found. Row ignored"))
        continue
    for col, allowed in (("Team", lists.get("department", [])), ("Status", lists["task_status"]), ("Priority", lists["priority"])):
        if not check(e[col], allowed, where, col):
            e[col] = None
    for col in ("Owner", "Waiting on"):
        if e[col] and e[col] not in people:
            report["skipped"].append((where, f"{col} '{e[col]}' is not in the people list. Saved as typed; add them in Settings > People"))
    if not rid:
        if not e["Task"] or not brand:
            report["skipped"].append((where, "New row needs a Brand and a Task. Row ignored"))
            continue
        next_n[brand] = next_n.get(brand, 0) + 1
        nid = f"{brand}-{next_n[brand]:03d}"
        blocker = {"text": e["Blocked because"] or "Reason not recorded", "person": e["Waiting on"] or e["Owner"]} if e["Status"] == "Blocked" else None
        queue("tasks", nid, {"brand": brand, "department": e["Team"] or "Other", "title": e["Task"], "owner": e["Owner"],
                             "support": split(e["Support team"]), "status": e["Status"] or "Not Started", "priority": e["Priority"] or "",
                             "start": e["Start"] or "", "deadline": e["Deadline"] or "", "notes": e["Notes / next step"],
                             "brief_link": e["Brief link"], "creative_link": e["Creative link"], "blocker": blocker,
                             "source": f"Excel import {datetime.date.today():%d %b %Y}"}, op="set")
        report["new"].append((where, f"New task for {brand_txt}: {e['Task']}"))
        continue
    seen.add(rid)
    t = live["tasks"].get(rid)
    if t is None:
        report["skipped"].append((where, f"Task {rid} no longer exists in the tracker"))
        continue
    blk_now = (e["Blocked because"], e["Waiting on"])
    apply("Tasks", "tasks", rid, r, e, {
        "Brand": ("brand", lambda s: by_name[s.lower()]), "Team": ("department", ident), "Task": ("title", ident), "Owner": ("owner", ident),
        "Support team": ("support", split), "Status": ("status", ident), "Priority": ("priority", ident), "Start": ("start", ident),
        "Deadline": ("deadline", ident), "Notes / next step": ("notes", ident), "Brief link": ("brief_link", ident),
        "Creative link": ("creative_link", ident)}, t, e["Task"] or t.get("title"))
    before = base[("Tasks", rid)][1]
    status_now = e["Status"] or t.get("status")
    if blk_now != (before.get("Blocked because", ""), before.get("Waiting on", "")) or (status_now == "Blocked" and e["Status"] != before.get("Status")):
        if status_now == "Blocked":
            if not blk_now[0] or not blk_now[1]:
                report["skipped"].append((where, f"{t['title']}: Blocked needs 'Blocked because' and 'Waiting on'. Filled what was missing"))
            queue("tasks", rid, {"blocker": {"text": blk_now[0] or "Reason not recorded", "person": blk_now[1] or e["Owner"] or t.get("owner", "")}})
            report["applied"].append((where, f"{t['title']}: blocker set"))
    if e["Status"] and e["Status"] != "Blocked" and before.get("Status") == "Blocked":
        queue("tasks", rid, {"blocker": None})
for (sh, rid), _ in base.items():
    if sh == "Tasks" and rid not in seen:
        report["missing"].append((f"Tasks {rid}", "Row missing from the file. Task kept in the tracker"))

# ---------------------------------------------------------------- Brands
for r, v in rows(wb["Brands"]):
    rid = txt(v.get("ID"))
    if rid not in live["brands"]:
        report["skipped"].append((f"Brands row {r}", "New brands are added in the tracker (Add brand). Row ignored"))
        continue
    e = {k: txt(v.get(k)) for k in ("Brand", "Founder", "Website", "Stage", "Brand lead", "Sprint", "Paused", "Next step")}
    if not check(e["Stage"], lists["stage"], f"Brands row {r}", "Stage"):
        e["Stage"] = None
    apply("Brands", "brands", rid, r, e, {"Brand": ("name", ident), "Founder": ("founder", ident), "Website": ("website", ident),
          "Stage": ("stage", ident), "Brand lead": ("lead", ident), "Sprint": ("sprint_label", ident), "Paused": ("paused", yes),
          "Next step": ("next_step", ident)}, live["brands"][rid], e["Brand"])

# ---------------------------------------------------------------- SOW
for r, v in rows(wb["SOW"]):
    rid = txt(v.get("ID"))
    if rid not in live["sow"]:
        report["skipped"].append((f"SOW row {r}", "Unknown SOW row. Ignored"))
        continue
    w = f"SOW row {r}"
    e = {"SOW stage": txt(v.get("SOW stage")), "Owner": txt(v.get("Owner")), "Support team": txt(v.get("Support team")),
         "Start date": date_val(v.get("Start date"), w), "Draft due": date_val(v.get("Draft due"), w), "Target sign": date_val(v.get("Target sign"), w),
         "Blocker / dependency": txt(v.get("Blocker / dependency")), "Next step": txt(v.get("Next step")), "SOW link": txt(v.get("SOW link")),
         "Growth Forecast link": txt(v.get("Growth Forecast link")), "GF updated": txt(v.get("GF updated")), "Batch": txt(v.get("Batch"))}
    if not check(e["SOW stage"], lists["sow_stage"], w, "SOW stage"):
        e["SOW stage"] = None
    apply("SOW", "sow", rid, r, e, {"SOW stage": ("stage", ident), "Owner": ("owner", ident), "Support team": ("support", split),
          "Start date": ("start_date", ident), "Draft due": ("draft_due", ident), "Target sign": ("target_sign", ident),
          "Blocker / dependency": ("blocker", ident), "Next step": ("next_step", ident), "SOW link": ("sow_link", ident),
          "Growth Forecast link": ("gf_link", ident), "GF updated": ("gf_updated", yes), "Batch": ("batch", ident)},
          live["sow"][rid], "SOW " + live["brands"].get(rid, {}).get("name", rid))

# ---------------------------------------------------------------- Platform access
for r, v in rows(wb["Platform Access"]):
    rid = txt(v.get("ID"))
    w = f"Platform Access row {r}"
    if not rid:
        report["skipped"].append((w, "New platform rows are added in the tracker for now. Row ignored"))
        continue
    if rid not in live["access"]:
        report["skipped"].append((w, "Unknown row. Ignored"))
        continue
    e = {"Platform": txt(v.get("Platform")), "Access level": txt(v.get("Access level")), "Needed for": txt(v.get("Needed for")),
         "Valerie owner": txt(v.get("Valerie owner")), "Brand contact": txt(v.get("Brand contact")),
         "Requested on": date_val(v.get("Requested on"), w), "Deadline": date_val(v.get("Deadline"), w), "Status": txt(v.get("Status")),
         "Blocker": txt(v.get("Blocker")), "Notes": txt(v.get("Notes"))}
    if not check(e["Status"], lists["access_status"], w, "Status"):
        e["Status"] = None
    apply("Platform Access", "access", rid, r, e, {"Platform": ("platform", ident), "Access level": ("access_level", ident),
          "Needed for": ("needed_for", ident), "Valerie owner": ("owner", ident), "Brand contact": ("brand_contact", ident),
          "Requested on": ("requested_date", ident), "Deadline": ("deadline", ident), "Status": ("status", ident),
          "Blocker": ("blocker", ident), "Notes": ("notes", ident)}, live["access"][rid], e["Platform"])

# ---------------------------------------------------------------- Onboarding steps grid
ws = wb["Onboarding Steps"]
head = [txt(c.value) for c in ws[1]]
for ci, name in enumerate(head[4:], 4):
    bid = by_name.get(name.lower())
    if not bid or ("Onboarding Steps", bid) not in base:
        continue
    ver0, before = base[("Onboarding Steps", bid)]
    live_s = (live["brand_steps"].get(bid) or {}).get("s") or {}
    for row in ws.iter_rows(min_row=2):
        code, now = txt(row[0].value), txt(row[ci].value)
        if not code or now == before.get(code, ""):
            continue
        w = f"Onboarding Steps {code} / {name}"
        if not check(now, lists["checklist"], w, "Status"):
            continue
        cur = live_s.get(code) or {}
        if VERS.get(f"brand_steps/{bid}", 0) != ver0 and cur.get("status", "") != before.get(code, ""):
            report["conflicts"].append((w, f"Tracker now says '{cur.get('status', '')}', your file says '{now}'. Kept the tracker value"))
            continue
        live_s[code] = dict(cur, status=now)          # later gate checks in this column see this change
        st = live["steps"].get(code) or {}
        if now == "Done" and st.get("is_gate"):
            open_ = [x["code"] for x in sorted(live["steps"].values(), key=lambda x: x["order"])
                     if x["order"] < st["order"] and not x.get("is_optional") and (live_s.get(x["code"]) or {}).get("status") not in ("Done", "N/A")]
            if open_:
                live_s[code] = cur
                report["skipped"].append((w, f"{code} is a gate. Finish {len(open_)} earlier steps first: {', '.join(open_[:4])}"))
                continue
        queue("brand_steps", bid, {"s": {code: dict(cur, status=now)}})
        report["applied"].append((w, f"'{before.get(code, '') or 'Not Started'}' to '{now}'"))

# ---------------------------------------------------------------- Best practices
secs = lists["bp_section"]
for r, v in rows(wb["Best Practices"]):
    rid = txt(v.get("ID"))
    if rid not in live["bps"]:
        report["skipped"].append((f"Best Practices row {r}", "Start tracking a brand in the tracker first. Row ignored"))
        continue
    w = f"Best Practices row {r}"
    e = {"BP owner": txt(v.get("BP owner")), "Due date": date_val(v.get("Due date"), w), "Blocker": txt(v.get("Blocker")),
         "BP document link": txt(v.get("BP document link")), "Notes": txt(v.get("Notes"))}
    m = {"BP owner": ("owner", ident), "Due date": ("due_date", ident), "Blocker": ("blocker", ident),
         "BP document link": ("doc_link", ident), "Notes": ("notes", ident)}
    apply("Best Practices", "bps", rid, r, e, m, live["bps"][rid], "BPs " + live["brands"][rid]["name"])
    ver0, before = base[("Best Practices", rid)]
    live_secs = live["bps"][rid].get("sections") or {}
    for s in secs:
        now = txt(v.get(s))
        if not now or now == before.get(s, "") or not check(now, lists["bp_status"], w, s):
            continue
        if VERS.get(f"bps/{rid}", 0) != ver0 and live_secs.get(s, "Not Started") != before.get(s, ""):
            report["conflicts"].append((w, f"{s}: tracker now says '{live_secs.get(s)}', your file says '{now}'. Kept the tracker value"))
            continue
        queue("bps", rid, {"sections": {s: now}})
        report["applied"].append((w, f"BPs {live['brands'][rid]['name']}: {s} '{before.get(s, '')}' to '{now}'"))

# ---------------------------------------------------------------- output
os.makedirs(os.path.join(OUT, "docs"), exist_ok=True)
batch = []
for (coll, rid), w in sorted(writes.items()):
    path = os.path.join(OUT, "docs", f"{coll}__{rid}.json")
    data = dict(w["data"])
    if w["op"] == "update":
        data["updated_at"] = int(datetime.datetime.now().timestamp() * 1000)
        data["updated_by"] = "excel-import"
    json.dump(data, open(path, "w"), ensure_ascii=False, indent=1)
    entry = {"op": w["op"], "collection": coll, "doc_id": rid, "file_path": path}
    if w["op"] == "update":
        entry["if_version"] = VERS.get(f"{coll}/{rid}", 0)
    batch.append(entry)
for i in range(0, len(batch), 50):
    json.dump(batch[i:i + 50], open(os.path.join(OUT, f"batch_{i // 50}.json"), "w"))
with open(os.path.join(OUT, "report.md"), "w") as f:
    f.write(f"# Excel import report\n\nFile: {os.path.basename(XLSX)}\n\n| Result | Count |\n| --- | --- |\n")
    for k, label in (("applied", "Cells applied"), ("new", "New tasks"), ("conflicts", "Conflicts (tracker kept)"),
                     ("skipped", "Skipped or fixed"), ("missing", "Rows missing from file")):
        f.write(f"| {label} | {len(report[k])} |\n")
    f.write(f"| Records to write | {len(batch)} |\n")
    for k, label in (("applied", "Applied"), ("new", "New tasks"), ("conflicts", "Conflicts"), ("skipped", "Skipped or fixed"), ("missing", "Missing rows")):
        if report[k]:
            f.write(f"\n## {label}\n\n| Where | What |\n| --- | --- |\n")
            for a, b in report[k]:
                f.write(f"| {a} | {str(b).replace('|', '/')} |\n")
print(json.dumps({k: len(v) for k, v in report.items()}), "records:", len(batch))

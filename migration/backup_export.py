"""Full Excel backup of the live tracker.

Usage:
    python3 -I migration/backup_export.py <live_dir> <out.xlsx>

Builds the editable workbook from sync_export.py (Tasks, Brands, SOW, Platform Access,
Onboarding Steps, Best Practices, Lists) and adds the records that workbook leaves out:
Campaigns, People, Escalations, Comments and the Change log, plus an About sheet with counts.
"""
import datetime
import json
import os
import subprocess
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

LIVE, OUT = sys.argv[1], sys.argv[2]
HERE = os.path.dirname(os.path.abspath(__file__))
if not os.path.exists(os.path.join(LIVE, "versions.json")):
    json.dump({}, open(os.path.join(LIVE, "versions.json"), "w"))
subprocess.run([sys.executable, "-I", os.path.join(HERE, "sync_export.py"), LIVE, OUT], check=True)


def load(coll):
    p = os.path.join(LIVE, coll)
    if not os.path.isdir(p):
        return {}
    return {f[:-5]: json.load(open(os.path.join(p, f))) for f in sorted(os.listdir(p)) if f.endswith(".json")}


def when(ms):
    try:
        return datetime.datetime.utcfromtimestamp(int(ms) / 1000).strftime("%Y-%m-%d %H:%M")
    except (TypeError, ValueError):
        return ""


HEAD = PatternFill("solid", fgColor="1B3A47")
wb = load_workbook(OUT)


def add(title, cols, widths, rows):
    ws = wb.create_sheet(title)
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for r in rows:
        ws.append(["" if v is None else v for v in r])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    return len(rows)


brands = load("brands")
bname = lambda b: (brands.get(b) or {}).get("name", b)
counts = {}

camps = load("campaigns")
counts["Campaigns"] = add("Campaigns", ["ID", "Brand", "Campaign", "Status", "Owner", "Start", "End", "What it is for", "Success looks like", "Link"],
                          [24, 16, 34, 12, 14, 12, 12, 60, 36, 30],
                          [[k, bname(c.get("brand")), c.get("name"), c.get("status"), c.get("owner"), c.get("start"), c.get("end"),
                            c.get("purpose"), c.get("goal"), c.get("link")] for k, c in camps.items()])

tasks = load("tasks")
cname = lambda c: (camps.get(c) or {}).get("name", c or "")
counts["Task extras"] = add("Task extras", ["ID", "Brand", "Task", "Campaign", "Board list", "Vertical", "Checklist", "Blocker since"],
                            [14, 16, 40, 30, 16, 14, 60, 14],
                            [[k, bname(t.get("brand")), t.get("title"), cname(t.get("campaign")), t.get("list", ""), t.get("vertical", ""),
                              "; ".join(("[x] " if i.get("done") else "[ ] ") + str(i.get("text", "")) for i in (t.get("checklist") or [])),
                              (t.get("blocker") or {}).get("at", "") if isinstance(t.get("blocker"), dict) else ""]
                             for k, t in tasks.items() if t.get("campaign") or t.get("list") or t.get("vertical") or t.get("checklist")])

people = (load("config").get("people") or {}).get("people", [])
counts["People"] = add("People", ["Name", "Active", "Leadership", "Kind", "Also known as"], [20, 9, 12, 10, 30],
                       [[p.get("name"), "Yes" if p.get("active") else "No", "Yes" if p.get("leadership") else "",
                         p.get("kind"), ", ".join(p.get("aliases") or [])] for p in people])

esc = load("escalations")
counts["Escalations"] = add("Escalations", ["ID", "Brand", "Item", "Asked", "Needed by", "Raised by", "Raised", "State", "Decision note", "Decided by"],
                            [20, 16, 34, 50, 12, 14, 16, 10, 40, 14],
                            [[k, bname(e.get("brand")), e.get("title"), e.get("ask"), e.get("due"), e.get("by"), when(e.get("at")),
                              e.get("state"), e.get("note"), e.get("decided_by")] for k, e in esc.items()])

com = load("comments")
counts["Comments"] = add("Comments", ["When", "Brand", "Task", "By", "Kind", "Text"], [16, 16, 16, 14, 10, 80],
                         [[when(c.get("at")), bname(c.get("brand")), c.get("task", ""), c.get("by"), c.get("kind", "comment"), c.get("text")]
                          for c in sorted(com.values(), key=lambda c: c.get("at") or 0)])

log_rows = []
for lid, day in load("log").items():
    for e in day.get("entries") or []:
        log_rows.append([when(e.get("at")), day.get("person") or "", e.get("what"), e.get("label") or e.get("id"), e.get("field"), e.get("old"), e.get("new")])
log_rows.sort(key=lambda r: r[0], reverse=True)
counts["Change log"] = add("Change log", ["When (UTC)", "Who", "Area", "Record", "Field", "Was", "Now"], [17, 14, 12, 34, 16, 40, 40], log_rows)

raw = {c: len(load(c)) for c in sorted(os.listdir(LIVE)) if os.path.isdir(os.path.join(LIVE, c))}
ab = wb.create_sheet("About this backup", 0)
ab.append(["Valerie COE Tracker backup"])
ab["A1"].font = Font(bold=True, size=14)
ab.append(["Taken", datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")])
ab.append(["What it is", "A copy of every record in the live tracker. Edit the white cells only if you plan to send it back for import."])
ab.append([])
ab.append(["Records in the tracker", "Count"])
for c, n in raw.items():
    ab.append([c, n])
ab.append([])
ab.append(["Extra sheets in this file", "Rows"])
for k, n in counts.items():
    ab.append([k, n])
ab.column_dimensions["A"].width = 28
ab.column_dimensions["B"].width = 90
wb.active = 0
wb.save(OUT)
print("backup", OUT, raw, counts)
